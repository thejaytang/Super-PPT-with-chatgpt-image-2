#!/usr/bin/env python3
"""Generate Super PPT visual proofs and editable PowerPoint decks."""

from __future__ import annotations

import argparse
import base64
import hashlib
import json
import sys
import time
from pathlib import Path
from typing import Any


REQUIRED_DECK_FIELDS = ("title",)
REQUIRED_SLIDE_FIELDS = ("number", "role", "title", "takeaway", "visual_intent")
SUPPORTED_ASPECT_RATIOS = {"16:9": (13.333333, 7.5), "4:3": (10.0, 7.5)}
SUPPORTED_GENERATION_MODES = {"hybrid_editable_text", "full_slide_image"}


class PlanError(ValueError):
    """Raised when the plan file is invalid."""


def load_plan(path: Path) -> dict[str, Any]:
    try:
        with path.open("r", encoding="utf-8") as handle:
            plan = json.load(handle)
    except json.JSONDecodeError as exc:
        raise PlanError(f"Invalid JSON: {exc}") from exc

    if not isinstance(plan, dict):
        raise PlanError("Plan must be a JSON object.")
    return plan


def validate_plan(plan: dict[str, Any]) -> None:
    deck = plan.get("deck")
    visual_system = plan.get("visual_system")
    slides = plan.get("slides")

    if not isinstance(deck, dict):
        raise PlanError("Missing object: deck")
    for field in REQUIRED_DECK_FIELDS:
        if not deck.get(field):
            raise PlanError(f"Missing deck.{field}")

    aspect_ratio = deck.get("aspect_ratio", "16:9")
    if aspect_ratio not in SUPPORTED_ASPECT_RATIOS:
        raise PlanError(f"Unsupported deck.aspect_ratio: {aspect_ratio}. Use one of {sorted(SUPPORTED_ASPECT_RATIOS)}")

    generation_mode = deck.get("generation_mode", "hybrid_editable_text")
    if generation_mode not in SUPPORTED_GENERATION_MODES:
        raise PlanError(f"Unsupported deck.generation_mode: {generation_mode}. Use one of {sorted(SUPPORTED_GENERATION_MODES)}")

    if not isinstance(visual_system, dict):
        raise PlanError("Missing object: visual_system")
    palette = visual_system.get("palette")
    if not isinstance(palette, dict) or not palette:
        raise PlanError("Missing visual_system.palette")
    for key, value in palette.items():
        if not isinstance(value, str) or not value.startswith("#") or len(value) not in {4, 7}:
            raise PlanError(f"visual_system.palette.{key} must be a hex color string")

    if not isinstance(slides, list) or not slides:
        raise PlanError("slides must be a non-empty list")

    seen_numbers: set[int] = set()
    for index, slide in enumerate(slides, start=1):
        if not isinstance(slide, dict):
            raise PlanError(f"slides[{index}] must be an object")
        for field in REQUIRED_SLIDE_FIELDS:
            if not slide.get(field):
                raise PlanError(f"Missing slides[{index}].{field}")

        number = slide["number"]
        if not isinstance(number, int):
            raise PlanError(f"slides[{index}].number must be an integer")
        if number in seen_numbers:
            raise PlanError(f"Duplicate slide number: {number}")
        seen_numbers.add(number)

        content = slide.get("content")
        if content is not None and not isinstance(content, (list, str, dict)):
            raise PlanError(f"slides[{index}].content must be a list, string, or object")
        if isinstance(content, list) and len(content) > 7:
            raise PlanError(f"slides[{index}].content has too many items for stable image text; keep it to 7 or fewer")

        editable_elements = slide.get("editable_elements", [])
        if editable_elements and not isinstance(editable_elements, list):
            raise PlanError(f"slides[{index}].editable_elements must be a list")


def format_value(value: Any) -> str:
    if isinstance(value, dict):
        return ", ".join(f"{key}: {item}" for key, item in value.items())
    if isinstance(value, list):
        return "; ".join(str(item) for item in value)
    return str(value)


def build_prompt(deck: dict[str, Any], visual_system: dict[str, Any], slide: dict[str, Any]) -> str:
    generation_mode = deck.get("generation_mode", "hybrid_editable_text")
    if slide.get("prompt"):
        base_prompt = str(slide["prompt"]).strip()
    else:
        base_prompt = (
            "Create a 16:9 full-slide PowerPoint visual proof for a polished business presentation.\n"
            f"Deck: {deck.get('title')}. Slide {slide['number']}: {slide['title']}. Role: {slide['role']}.\n"
            f"Takeaway: {slide['takeaway']}.\n"
            f"Visual intent: {slide['visual_intent']}.\n"
            f"Content cues: {format_value(slide.get('content', []))}.\n"
        )

    system_lines = [
        f"Language: {deck.get('language', 'match the source content')}.",
        f"Aspect ratio: {deck.get('aspect_ratio', '16:9')}.",
        f"Visual system style: {visual_system.get('style_name', 'consistent professional presentation')}.",
        f"Palette: {format_value(visual_system.get('palette', {}))}.",
        f"Typography: {visual_system.get('typography', 'modern sans-serif with clear hierarchy')}.",
        f"UI motif: {visual_system.get('ui_motif', 'clean presentation UI')}.",
        f"Composition: {visual_system.get('composition', 'consistent grid and margins')}.",
    ]
    if visual_system.get("image_style"):
        system_lines.append(f"Image style: {visual_system['image_style']}.")
    if visual_system.get("negative_rules"):
        system_lines.append(f"Negative rules: {format_value(visual_system['negative_rules'])}.")

    if generation_mode == "hybrid_editable_text":
        system_lines.append(
            "Phase: visual proof. Create a pure image visual design that reserves clean zones for later editable "
            "PowerPoint text overlays. Do not render exact final titles, long bullets, detailed numbers, citations, "
            "or dense tables inside the image. Use subtle placeholder blocks or label areas when needed."
        )
    else:
        system_lines.append(
            "Phase: visual proof and final raster slide. Render only concise, readable text that is explicitly requested."
        )

    system_lines.append(
        "Keep the slide consistent with the same deck. Avoid unrelated colors, random logos, watermarks, "
        "decorative gradient blobs, tiny unreadable text, and inconsistent icon styles."
    )

    return base_prompt + "\n" + "\n".join(system_lines)


def collect_prompts(plan: dict[str, Any]) -> list[tuple[dict[str, Any], str]]:
    deck = plan["deck"]
    visual_system = plan["visual_system"]
    return [(slide, build_prompt(deck, visual_system, slide)) for slide in sorted(plan["slides"], key=lambda item: item["number"])]


def write_prompt_markdown(prompts: list[tuple[dict[str, Any], str]], path: Path) -> None:
    lines: list[str] = ["# Super PPT Image Prompts", ""]
    for slide, prompt in prompts:
        lines.extend([f"## Slide {slide['number']}: {slide['title']}", "", "```text", prompt, "```", ""])
    path.write_text("\n".join(lines), encoding="utf-8")


def prompt_hash(prompt: str) -> str:
    return hashlib.sha256(prompt.encode("utf-8")).hexdigest()[:16]


def generate_image(prompt: str, output_path: Path, model: str, size: str, quality: str, max_retries: int) -> None:
    try:
        from openai import OpenAI
    except ImportError as exc:
        raise RuntimeError("Missing dependency: install openai to generate images.") from exc

    client = OpenAI()
    last_error: Exception | None = None
    for attempt in range(1, max_retries + 1):
        try:
            response = client.images.generate(
                model=model,
                prompt=prompt,
                size=size,
                quality=quality,
                n=1,
            )
            break
        except Exception as exc:  # SDK exceptions vary by version.
            last_error = exc
            if attempt == max_retries:
                raise RuntimeError(f"Image generation failed after {max_retries} attempts: {exc}") from exc
            time.sleep(min(2**attempt, 10))
    else:
        raise RuntimeError(f"Image generation failed: {last_error}")

    image = response.data[0]
    b64_json = getattr(image, "b64_json", None)
    if not b64_json:
        raise RuntimeError("Image response did not include b64_json. Check SDK/model output settings.")
    output_path.write_bytes(base64.b64decode(b64_json))


def ensure_png_or_jpeg(path: Path) -> None:
    suffix = path.suffix.lower()
    if suffix not in {".png", ".jpg", ".jpeg"}:
        raise RuntimeError(f"PowerPoint image must be PNG or JPEG: {path}")


def hex_to_rgb(value: str) -> tuple[int, int, int]:
    value = value.lstrip("#")
    if len(value) == 3:
        value = "".join(char * 2 for char in value)
    return int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16)


def default_editable_elements(slide_data: dict[str, Any]) -> list[dict[str, Any]]:
    elements: list[dict[str, Any]] = [
        {
            "type": "text",
            "text": slide_data["title"],
            "x": 0.7,
            "y": 0.45,
            "w": 8.8,
            "h": 0.55,
            "font_size": 28,
            "bold": True,
            "panel": False,
        },
        {
            "type": "text",
            "text": slide_data["takeaway"],
            "x": 0.72,
            "y": 1.08,
            "w": 9.8,
            "h": 0.38,
            "font_size": 14,
            "bold": False,
            "panel": False,
        },
    ]

    content = slide_data.get("content")
    if isinstance(content, list) and content:
        elements.append(
            {
                "type": "bullet_list",
                "items": [str(item) for item in content[:5]],
                "x": 0.9,
                "y": 5.15,
                "w": 5.8,
                "h": 1.45,
                "font_size": 13,
                "panel": True,
            }
        )
    return elements


def add_text_element(slide: Any, element: dict[str, Any], defaults: dict[str, str]) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    x = Inches(float(element.get("x", 0.7)))
    y = Inches(float(element.get("y", 0.5)))
    w = Inches(float(element.get("w", 6.0)))
    h = Inches(float(element.get("h", 0.5)))

    if element.get("panel"):
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        panel.fill.solid()
        panel.fill.fore_color.rgb = RGBColor(*hex_to_rgb(defaults["surface"]))
        panel.line.color.rgb = RGBColor(*hex_to_rgb(defaults["muted"]))

    textbox = slide.shapes.add_textbox(x, y, w, h)
    frame = textbox.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.05)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)

    element_type = element.get("type", "text")
    if element_type == "bullet_list":
        items = element.get("items", [])
        if isinstance(items, str):
            items = [items]
        for index, item in enumerate(items):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = str(item)
            paragraph.level = 0
            paragraph.font.size = Pt(float(element.get("font_size", 13)))
            paragraph.font.color.rgb = RGBColor(*hex_to_rgb(element.get("color", defaults["text"])))
            paragraph.font.bold = bool(element.get("bold", False))
            paragraph.space_after = Pt(4)
    else:
        paragraph = frame.paragraphs[0]
        paragraph.text = str(element.get("text", ""))
        paragraph.font.size = Pt(float(element.get("font_size", 16)))
        paragraph.font.color.rgb = RGBColor(*hex_to_rgb(element.get("color", defaults["text"])))
        paragraph.font.bold = bool(element.get("bold", False))

    align = str(element.get("align", "left")).lower()
    if align == "center":
        frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    elif align == "right":
        frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    else:
        frame.paragraphs[0].alignment = PP_ALIGN.LEFT


def assemble_pptx(plan: dict[str, Any], image_paths: list[Path], output_path: Path, editable: bool = False) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise RuntimeError("Missing dependency: install python-pptx to assemble the deck.") from exc

    presentation = Presentation()
    width, height = SUPPORTED_ASPECT_RATIOS[plan["deck"].get("aspect_ratio", "16:9")]
    presentation.slide_width = Inches(width)
    presentation.slide_height = Inches(height)
    blank_layout = presentation.slide_layouts[6]

    palette = plan["visual_system"].get("palette", {})
    defaults = {
        "surface": palette.get("surface", "#FFFFFF"),
        "text": palette.get("text", "#111827"),
        "muted": palette.get("muted", "#D0D5DD"),
    }

    slides = sorted(plan["slides"], key=lambda item: item["number"])
    for slide_data, image_path in zip(slides, image_paths):
        ensure_png_or_jpeg(image_path)
        slide = presentation.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            str(image_path),
            0,
            0,
            width=presentation.slide_width,
            height=presentation.slide_height,
        )
        slide.name = f"{slide_data['number']:02d} {slide_data['title']}"
        if editable and plan["deck"].get("generation_mode", "hybrid_editable_text") == "hybrid_editable_text":
            elements = slide_data.get("editable_elements") or default_editable_elements(slide_data)
            for element in elements:
                add_text_element(slide, element, defaults)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)


def write_manifest(
    plan: dict[str, Any],
    prompts: list[tuple[dict[str, Any], str]],
    image_paths: list[Path],
    manifest_path: Path,
    visual_output_path: Path,
    editable_output_path: Path | None,
    model: str,
    size: str,
    quality: str,
) -> None:
    slides = []
    for slide, prompt in prompts:
        image_path = next((path for path in image_paths if path.stem == f"slide-{slide['number']:02d}"), None)
        slides.append(
            {
                "number": slide["number"],
                "title": slide["title"],
                "role": slide["role"],
                "takeaway": slide["takeaway"],
                "qa_notes": slide.get("qa_notes", []),
                "prompt_hash": prompt_hash(prompt),
                "image_path": str(image_path) if image_path else None,
                "editable_element_count": len(slide.get("editable_elements", [])),
            }
        )

    manifest = {
        "deck_title": plan["deck"]["title"],
        "aspect_ratio": plan["deck"].get("aspect_ratio", "16:9"),
        "generation_mode": plan["deck"].get("generation_mode", "hybrid_editable_text"),
        "model": model,
        "size": size,
        "quality": quality,
        "visual_output": str(visual_output_path),
        "editable_output": str(editable_output_path) if editable_output_path else None,
        "slides": slides,
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate Super PPT visual proofs and editable PPTX decks.")
    parser.add_argument("plan", type=Path, help="Path to plan.json")
    parser.add_argument("--output", type=Path, default=Path("super-ppt-visual.pptx"), help="Visual proof .pptx path")
    parser.add_argument("--editable-output", type=Path, default=Path("super-ppt-editable.pptx"), help="Editable .pptx path")
    parser.add_argument("--image-dir", type=Path, default=Path("super-ppt-images"), help="Directory for slide images")
    parser.add_argument("--model", default="gpt-image-2", help="OpenAI image model to use")
    parser.add_argument("--size", default="1536x1024", help="Image API size parameter")
    parser.add_argument("--quality", default="high", help="Image quality parameter")
    parser.add_argument("--dry-run", action="store_true", help="Validate the plan and write prompts without API calls")
    parser.add_argument("--prompt-out", type=Path, default=Path("super-ppt-prompts.md"), help="Prompt markdown output")
    parser.add_argument("--manifest-out", type=Path, default=Path("super-ppt-manifest.json"), help="Generation manifest output")
    parser.add_argument("--reuse-images", action="store_true", help="Assemble PPTX from existing images in --image-dir")
    parser.add_argument("--skip-editable", action="store_true", help="Only assemble the visual proof deck")
    parser.add_argument("--max-retries", type=int, default=3, help="Maximum attempts per generated image")
    return parser.parse_args(argv)


def main(argv: list[str]) -> int:
    args = parse_args(argv)

    try:
        plan = load_plan(args.plan)
        validate_plan(plan)
        prompts = collect_prompts(plan)
        write_prompt_markdown(prompts, args.prompt_out)

        if args.dry_run:
            print(f"Validated {len(prompts)} slides. Wrote prompts to {args.prompt_out}.")
            return 0

        args.image_dir.mkdir(parents=True, exist_ok=True)
        image_paths: list[Path] = []
        for slide, prompt in prompts:
            image_path = args.image_dir / f"slide-{slide['number']:02d}.png"
            if args.reuse_images and image_path.exists():
                print(f"Reusing {image_path}")
            else:
                print(f"Generating {image_path} with {args.model}")
                generate_image(prompt, image_path, args.model, args.size, args.quality, args.max_retries)
            image_paths.append(image_path)

        assemble_pptx(plan, image_paths, args.output, editable=False)
        editable_output_path = None
        if not args.skip_editable:
            editable_output_path = args.editable_output
            assemble_pptx(plan, image_paths, editable_output_path, editable=True)
            print(f"Wrote {editable_output_path}")

        write_manifest(
            plan,
            prompts,
            image_paths,
            args.manifest_out,
            args.output,
            editable_output_path,
            args.model,
            args.size,
            args.quality,
        )
        print(f"Wrote {args.output}")
        print(f"Wrote {args.manifest_out}")
        return 0
    except (PlanError, RuntimeError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
